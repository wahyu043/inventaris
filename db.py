import sqlite3
import click
from flask import g
from werkzeug.security import generate_password_hash


def get_db():
    # g = objek penyimpanan sementara selama 1 request
    if "db" not in g:
        g.db = sqlite3.connect("database.db")
        g.db.row_factory = sqlite3.Row   # hasil query bisa diakses pakai nama kolom

        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")   # SQLite default-nya MATI, harus dinyalakan tiap koneksi
    return g.db


def close_db(e=None):
    conn = g.pop("db", None)
    if conn is not None:
        conn.close()


@click.command("init-db")
def init_db_command():
    """Buat tabel + user admin awal."""
    conn = get_db()
    with open("schema.sql") as f:
        conn.executescript(f.read())
    conn.execute(
        "INSERT OR IGNORE INTO users (username, password_hash) VALUES (?, ?)",
        ("admin", generate_password_hash("admin123", method="pbkdf2:sha256")),
    )
    conn.commit()
    click.echo("Database siap. Login: admin / admin123")

@click.command("import-excel")
@click.argument("path")
def import_excel_command(path):
    """Impor data transaksi dari file Excel."""
    from openpyxl import load_workbook

    conn = get_db()
    if conn.execute("SELECT COUNT(*) FROM transaksi").fetchone()[0] > 0:
        click.echo("Tabel transaksi sudah berisi data. Impor dibatalkan.")
        return

    ws = load_workbook(path, data_only=True).active
    jumlah_baris = 0
    for row in ws.iter_rows(min_row=4, values_only=True):   # baris 1-3 = judul + header
        no, tanggal, jenis, kode, jumlah, nama, faktur, stok = row
        if kode is None:
            continue
        # barang cukup dimasukkan sekali; kalau sudah ada, dilewati
        conn.execute("INSERT OR IGNORE INTO barang (kode, nama) VALUES (?, ?)", (kode, nama))
        conn.execute(
            "INSERT INTO transaksi (tanggal, jenis, kode_barang, jumlah, no_faktur) VALUES (?, ?, ?, ?, ?)",
            (tanggal.strftime("%Y-%m-%d"), jenis, kode, int(jumlah), faktur),
        )
        jumlah_baris += 1
    conn.commit()
    click.echo(f"{jumlah_baris} transaksi diimpor.")

def init_app(app):
    app.teardown_appcontext(close_db)   # tutup koneksi otomatis tiap request selesai
    app.cli.add_command(init_db_command)
    app.cli.add_command(import_excel_command)